from datetime import date, datetime, timezone
from io import BytesIO
import openpyxl
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from app.holders.models import Holder
from app.lifecycle.service import apply_event
from app.lifecycle.state_machine import LifecycleError
from app.masters.models import AssetCategory, AssetSubcategory, Company, CostCenter
from app.numbering.service import generate_code, get_active_rule
from app.assets.models import Asset

TEMPLATE_COLUMNS = [
    "legacy_asset_code", "company_code", "cost_center_code", "category_code",
    "subcategory_code", "description", "purchase_date", "holder_emp_code",
]


def build_template() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(TEMPLATE_COLUMNS)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def _lookup(session: AsyncSession, model, **filters):
    stmt = select(model)
    for key, value in filters.items():
        stmt = stmt.where(getattr(model, key) == value)
    return (await session.execute(stmt)).scalars().first()


def _parse_purchase_date(value) -> date:
    """Parses/validates the workbook's raw purchase_date cell value into a `date`.

    Shared by `_validate_rows` (so both preview and commit apply the exact same rule)
    rather than left for `commit_import` to parse on its own -- a malformed date used to
    pass preview as "valid" and then raise an uncaught ValueError inside commit_import's
    loop, silently 500-ing a request that preview had just told the caller was clean.
    """
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return datetime.strptime(value, "%Y-%m-%d").date()
        except ValueError:
            raise ValueError(f"purchase_date '{value}' is not a valid date (expected YYYY-MM-DD)")
    raise ValueError(f"purchase_date '{value}' is not a valid date (expected YYYY-MM-DD)")


async def _validate_rows(session: AsyncSession, content: bytes) -> tuple[list[dict], list[dict]]:
    wb = openpyxl.load_workbook(BytesIO(content))
    ws = wb.active
    valid_rows: list[dict] = []
    errors: list[dict] = []

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(v is None for v in row):
            continue
        legacy_code, company_code, cc_code, cat_code, sub_code, description, purchase_date, holder_code = row

        company = await _lookup(session, Company, code=company_code)
        if company is None:
            errors.append({"row": row_idx, "message": f"unknown company_code '{company_code}'"})
            continue
        cost_center = await _lookup(session, CostCenter, company_id=company.id, code=cc_code)
        if cost_center is None:
            errors.append({"row": row_idx, "message": f"unknown cost_center_code '{cc_code}'"})
            continue
        category = await _lookup(session, AssetCategory, code=cat_code)
        if category is None:
            errors.append({"row": row_idx, "message": f"unknown category_code '{cat_code}'"})
            continue
        subcategory = await _lookup(session, AssetSubcategory, category_id=category.id, code=sub_code)
        if subcategory is None:
            errors.append({"row": row_idx, "message": f"unknown subcategory_code '{sub_code}' for category '{cat_code}'"})
            continue
        holder = await _lookup(session, Holder, company_id=company.id, emp_code=holder_code)
        if holder is None:
            errors.append({"row": row_idx, "message": f"unknown holder_emp_code '{holder_code}'"})
            continue
        if not description:
            errors.append({"row": row_idx, "message": "description is required"})
            continue
        try:
            purchase_date = _parse_purchase_date(purchase_date)
        except ValueError as exc:
            errors.append({"row": row_idx, "message": str(exc)})
            continue

        valid_rows.append({
            "row": row_idx, "legacy_asset_code": legacy_code, "company": company, "cost_center": cost_center,
            "category": category, "subcategory": subcategory, "description": description,
            "purchase_date": purchase_date, "holder": holder,
        })

    return valid_rows, errors


async def preview_import(session: AsyncSession, content: bytes) -> dict:
    valid_rows, errors = await _validate_rows(session, content)
    return {
        "valid_rows": [{"row": r["row"], "legacy_asset_code": r["legacy_asset_code"], "description": r["description"]} for r in valid_rows],
        "errors": errors,
    }


async def commit_import(session: AsyncSession, content: bytes, actor: Holder) -> dict:
    valid_rows, errors = await _validate_rows(session, content)
    rule = await get_active_rule(session, valid_rows[0]["company"].id) if valid_rows else None

    imported = 0
    for r in valid_rows:
        tokens = {
            "cost_center.code": r["cost_center"].code, "category.code": r["category"].code,
            "subcategory.code": r["subcategory"].code, "company.code": r["company"].code, "location.code": "",
            "yyyy": "", "yy": "", "mm": "",
        }
        purchase_date = r["purchase_date"]
        try:
            # Each row's insert + ledger event runs inside its own SAVEPOINT: if
            # apply_event raises LifecycleError (e.g. a future-dated row, or any other
            # transition the state machine rejects), rolling back just this savepoint
            # discards that row's partially-flushed Asset insert (and the code-counter
            # increment from generate_code, which ran inside the same savepoint) without
            # aborting the whole commit -- matching the per-item try/except pattern
            # app.assets.router.bulk_move already uses for apply_event failures, so one
            # bad row degrades to a reported error instead of a bare 500 for the entire
            # batch.
            async with session.begin_nested():
                code = await generate_code(session, rule, tokens)
                asset = Asset(
                    asset_code=code, legacy_asset_code=r["legacy_asset_code"], company_id=r["company"].id,
                    cost_center_id=r["cost_center"].id, category_id=r["category"].id, subcategory_id=r["subcategory"].id,
                    description=r["description"], purchase_date=purchase_date,
                    # Initial status set directly here, not through apply_event -- this is
                    # the same documented exception used by app.assets.service.procure_assets:
                    # a freshly inserted row needs a non-null status/holder before the state
                    # machine has anything to transition from. It is always IN_STOCK here
                    # (never derived from the target holder's type) because apply_event's
                    # IMPORTED transition is only defined FROM IN_STOCK -- apply_event,
                    # called immediately below, derives and writes the real post-import
                    # status from the holder's type, exactly like procure_assets does for
                    # PROCURED. Precomputing e.g. ALLOTTED/INSTALLED here would make every
                    # non-IT_STOCK row's IMPORTED transition invalid before it even runs.
                    status="IN_STOCK",
                    current_holder_id=r["holder"].id, status_since=purchase_date,
                    created_by=actor.id, updated_by=actor.id,
                )
                session.add(asset)
                await session.flush()
                await apply_event(
                    session, asset, "IMPORTED", to_holder_id=r["holder"].id, actor=actor,
                    event_date=datetime.combine(purchase_date, datetime.min.time()).replace(tzinfo=timezone.utc),
                    remarks=f"Imported from legacy code {r['legacy_asset_code']}",
                )
            imported += 1
        except LifecycleError as exc:
            errors.append({"row": r["row"], "message": str(exc)})

    await session.flush()
    return {"imported": imported, "errors": errors}
