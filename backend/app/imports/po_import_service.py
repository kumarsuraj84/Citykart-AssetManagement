"""Imports a Purchase Order and its line items from Excel -- one row per
line item, with the PO's own header fields (Company Code/PO Number/PO
Date/Vendor Code/Cost Centre Code) repeated on every row belonging to that
PO, the same convention a real prepared spreadsheet already uses. Rows are
grouped by (Company, PO Number); every row sharing one PO Number becomes
one PurchaseOrder (via app.purchase_orders.service.create_purchase_order)
with that many PendingAsset lines (via add_pending_asset_line) -- ordinary
PENDING lines, exactly as if entered one at a time through "+ Add Line".
Nothing is auto-delivered; the imported PO still goes through the normal
Mark Delivery Done flow afterward.

Primary-Owner-only (see app.imports.router) -- kept consistent with every
other bulk-import capability, not because this one creates real Assets
directly (it doesn't; only staging PendingAsset rows).

Reuses app.imports.asset_import_service's own row-parsing helpers and
exceptions rather than reimplementing them."""
from io import BytesIO
import openpyxl
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from app.asset_users.models import AssetUser
from app.imports.asset_import_service import (
    ImportScopeError, ImportTemplateError,
    _is_blank, _lookup, _out_of_scope_message, _parse_date_value, _parse_optional_number,
    _parse_quantity, _parse_warranty_years, _read_header,
)
from app.masters.models import AssetCategory, AssetSubcategory, Brand, Company, CostCenter, Vendor
from app.purchase_orders.models import PurchaseOrder
from app.purchase_orders.service import add_pending_asset_line, create_purchase_order

PO_TEMPLATE_COLUMNS = [
    "Company Code", "PO Number", "PO Date", "Vendor Code", "Cost Centre Code",
    "Description", "Barcode", "Category Code", "Subcategory Code",
    "Brand Code", "Model", "Warranty Years", "Purchase Cost", "Tax %", "Quantity",
]
# Vendor Code/Subcategory Code/Brand Code/Model/Warranty Years/Purchase
# Cost/Tax %/Quantity are optional, matching "+ Add Line"'s own optional
# fields (add_pending_asset_line). Barcode is required here even though
# the service layer itself treats it as optional -- the Add Line dialog
# already makes it mandatory in the UI, so an import should hold the same
# real business rule.
PO_REQUIRED_COLUMNS = [
    "Company Code", "PO Number", "PO Date", "Cost Centre Code",
    "Description", "Barcode", "Category Code",
]


def build_template() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(PO_TEMPLATE_COLUMNS)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def _validate_rows(
    session: AsyncSession, content: bytes, allowed_company_ids: list[int] | None = None,
) -> tuple[list[dict], list[dict], list[dict], dict[tuple[int, str], list[dict]]]:
    """Returns (valid_rows, errors, scope_violations, groups). `groups` maps
    (company_id, po_number) -> the valid rows belonging to that PO, in file
    order -- the exact unit commit_import iterates to create one
    PurchaseOrder per group."""
    wb = openpyxl.load_workbook(BytesIO(content))
    ws = wb.active
    header = _read_header(ws)

    missing = [c for c in PO_REQUIRED_COLUMNS if c not in header]
    if missing:
        raise ImportTemplateError(f"missing required column(s): {', '.join(missing)} -- re-download the template")

    valid_rows: list[dict] = []
    errors: list[dict] = []
    scope_violations: list[dict] = []
    groups: dict[tuple[int, str], list[dict]] = {}
    # The first row's own header-field values per group, so every later row
    # sharing the same PO Number can be checked for agreement with them --
    # a PO can only ever have one PO Date/Vendor/Cost Centre.
    group_header_values: dict[tuple[int, str], dict] = {}

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(v is None for v in row):
            continue

        def cell(name: str):
            idx = header.get(name)
            return row[idx] if idx is not None and idx < len(row) else None

        company_code = cell("Company Code")
        company = await _lookup(session, Company, code=company_code)
        if company is None:
            errors.append({"row": row_idx, "field": "Company Code", "message": f"unknown Company Code '{company_code}'"})
            continue
        if allowed_company_ids is not None and company.id not in allowed_company_ids:
            violation = {"row": row_idx, "field": "Company Code", "message": _out_of_scope_message(company_code)}
            errors.append(violation)
            scope_violations.append(violation)
            continue

        po_number_raw = cell("PO Number")
        if _is_blank(po_number_raw):
            errors.append({"row": row_idx, "field": "PO Number", "message": "PO Number is required"})
            continue
        po_number = str(po_number_raw).strip()

        try:
            po_date = _parse_date_value(cell("PO Date"), "PO Date")
        except ValueError as exc:
            errors.append({"row": row_idx, "field": "PO Date", "message": str(exc)})
            continue

        vendor_code = cell("Vendor Code")
        vendor = None
        if not _is_blank(vendor_code):
            vendor = await _lookup(session, Vendor, code=vendor_code)
            if vendor is None:
                errors.append({"row": row_idx, "field": "Vendor Code", "message": f"unknown Vendor Code '{vendor_code}'"})
                continue

        cc_code = cell("Cost Centre Code")
        cost_center = await _lookup(session, CostCenter, company_id=company.id, code=cc_code)
        if cost_center is None:
            errors.append({"row": row_idx, "field": "Cost Centre Code", "message": f"unknown Cost Centre Code '{cc_code}'"})
            continue

        description = cell("Description")
        if _is_blank(description):
            errors.append({"row": row_idx, "field": "Description", "message": "Description is required"})
            continue

        barcode = cell("Barcode")
        if _is_blank(barcode):
            errors.append({"row": row_idx, "field": "Barcode", "message": "Barcode is required"})
            continue

        cat_code = cell("Category Code")
        category = await _lookup(session, AssetCategory, code=cat_code)
        if category is None:
            errors.append({"row": row_idx, "field": "Category Code", "message": f"unknown Category Code '{cat_code}'"})
            continue

        subcategory = None
        sub_code = cell("Subcategory Code")
        if not _is_blank(sub_code):
            subcategory = await _lookup(session, AssetSubcategory, category_id=category.id, code=sub_code)
            if subcategory is None:
                errors.append({
                    "row": row_idx, "field": "Subcategory Code",
                    "message": f"unknown Subcategory Code '{sub_code}' for category '{cat_code}'",
                })
                continue

        brand = None
        brand_code = cell("Brand Code")
        if not _is_blank(brand_code):
            brand = await _lookup(session, Brand, code=brand_code)
            if brand is None:
                errors.append({"row": row_idx, "field": "Brand Code", "message": f"unknown Brand Code '{brand_code}'"})
                continue

        try:
            warranty_years = _parse_warranty_years(cell("Warranty Years"))
            purchase_cost = _parse_optional_number(cell("Purchase Cost"), "Purchase Cost")
            tax_percent = _parse_optional_number(cell("Tax %"), "Tax %")
            quantity = _parse_quantity(cell("Quantity"))
        except ValueError as exc:
            errors.append({"row": row_idx, "message": str(exc)})
            continue

        key = (company.id, po_number)
        header_values = {"po_date": po_date, "vendor_id": vendor.id if vendor else None, "cost_center_id": cost_center.id}
        if key in group_header_values:
            if group_header_values[key] != header_values:
                errors.append({
                    "row": row_idx, "field": "PO Number",
                    "message": f"PO Number '{po_number}' rows disagree on PO Date/Vendor Code/Cost Centre Code -- every row for the same PO must match",
                })
                continue
        else:
            group_header_values[key] = header_values

        row_data = {
            "row": row_idx, "company": company, "po_number": po_number, "po_date": po_date,
            "vendor_id": vendor.id if vendor else None, "cost_center_id": cost_center.id,
            "cost_center_name": cost_center.name,
            "description": str(description).strip(), "barcode": str(barcode).strip(),
            "category_id": category.id, "category_name": category.name,
            "subcategory_id": subcategory.id if subcategory else None,
            "brand_id": brand.id if brand else None,
            "model": None if _is_blank(cell("Model")) else str(cell("Model")).strip(),
            "warranty_years": warranty_years, "purchase_cost": purchase_cost, "tax_percent": tax_percent,
            "quantity": quantity,
        }
        groups.setdefault(key, []).append(row_data)
        valid_rows.append(row_data)

    return valid_rows, errors, scope_violations, groups


async def _reject_existing_po_numbers(
    session: AsyncSession, valid_rows: list[dict], errors: list[dict], groups: dict[tuple[int, str], list[dict]],
) -> tuple[list[dict], list[dict], dict[tuple[int, str], list[dict]]]:
    """A PO Number has no database uniqueness constraint (a real,
    pre-existing gap in ordinary manual PO creation too -- see
    docs/ai/DECISIONS.md), so an import can't rely on a DB-level conflict
    to catch a collision. Checked here instead: importing a PO Number that
    already exists for that company is refused outright (every row in that
    group becomes an error), rather than silently grafting new lines onto
    a real existing PO."""
    bad_keys: set[tuple[int, str]] = set()
    for company_id, po_number in groups:
        existing = (await session.execute(
            select(PurchaseOrder).where(PurchaseOrder.company_id == company_id, PurchaseOrder.po_number == po_number)
        )).scalars().first()
        if existing is not None:
            bad_keys.add((company_id, po_number))
    if not bad_keys:
        return valid_rows, errors, groups

    kept_valid = []
    for r in valid_rows:
        key = (r["company"].id, r["po_number"])
        if key in bad_keys:
            errors.append({
                "row": r["row"], "field": "PO Number",
                "message": f"PO Number '{r['po_number']}' already exists for this company",
            })
        else:
            kept_valid.append(r)
    kept_groups = {k: v for k, v in groups.items() if k not in bad_keys}
    return kept_valid, errors, kept_groups


async def preview_import(session: AsyncSession, content: bytes, allowed_company_ids: list[int] | None = None) -> dict:
    valid_rows, errors, _scope_violations, groups = await _validate_rows(session, content, allowed_company_ids)
    valid_rows, errors, groups = await _reject_existing_po_numbers(session, valid_rows, errors, groups)
    preview_rows = [
        {
            "row": r["row"], "po_number": r["po_number"], "po_date": r["po_date"].isoformat(),
            "company": r["company"].name, "cost_center": r["cost_center_name"],
            "description": r["description"], "barcode": r["barcode"],
            "category": r["category_name"], "quantity": r["quantity"],
        }
        for r in valid_rows
    ]
    return {"valid_rows": preview_rows, "errors": errors}


async def commit_import(
    session: AsyncSession, content: bytes, actor: AssetUser, allowed_company_ids: list[int] | None = None,
) -> dict:
    valid_rows, errors, scope_violations, groups = await _validate_rows(session, content, allowed_company_ids)
    if scope_violations:
        raise ImportScopeError(f"{len(scope_violations)} row(s) target a company outside your scope -- nothing was imported")
    _valid_rows, errors, groups = await _reject_existing_po_numbers(session, valid_rows, errors, groups)

    pos_created = 0
    lines_created = 0
    for (company_id, po_number), rows in groups.items():
        first = rows[0]
        po = await create_purchase_order(session, {
            "company_id": company_id, "po_number": po_number, "po_date": first["po_date"],
            "vendor_id": first["vendor_id"], "cost_center_id": first["cost_center_id"],
        }, actor)
        pos_created += 1
        for r in rows:
            lines = await add_pending_asset_line(session, po, {
                "description": r["description"], "barcode": r["barcode"], "category_id": r["category_id"],
                "subcategory_id": r["subcategory_id"], "brand_id": r["brand_id"], "model": r["model"],
                "warranty_years": r["warranty_years"], "purchase_cost": r["purchase_cost"],
                "tax_percent": r["tax_percent"], "quantity": r["quantity"],
            }, actor)
            lines_created += len(lines)

    await session.flush()
    return {"pos_created": pos_created, "lines_created": lines_created, "errors": errors}
