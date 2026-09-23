from datetime import date, timedelta
from io import BytesIO

import openpyxl
from app.core.db import SessionLocal
from app.core.security import hash_password
from app.assets.service import procure_assets
from app.masters.models import Company, CostCenter, AssetCategory, AssetSubcategory, Location, Department
from app.holders.models import Holder
from app.numbering.models import CodeRule


async def _setup_company(session, suffix: str):
    co = Company(code=f"CKS-{suffix}", name=f"Export Test Co {suffix}")
    cat = AssetCategory(code=f"IT-{suffix}", name="IT")
    session.add_all([co, cat])
    await session.flush()
    sub = AssetSubcategory(category_id=cat.id, code="LAP", name="Laptop")
    cc = CostCenter(company_id=co.id, code="HO01", name="HO")
    loc = Location(code=f"HO-{suffix}", name="HO")
    dept = Department(name=f"IT-{suffix}")
    session.add_all([sub, cc, loc, dept])
    await session.flush()
    stock = Holder(company_id=co.id, emp_code=f"ITSTOCK-{suffix}", name=f"IT Stock-{suffix}", holder_type="IT_STOCK",
                    location_id=loc.id, department_id=dept.id, role="HOLDER")
    admin = Holder(company_id=co.id, emp_code=f"ITA-{suffix}", name=f"IT Admin {suffix}", holder_type="EMPLOYEE",
                    location_id=loc.id, department_id=dept.id, role="ADMIN",
                    password_hash=hash_password("Passw0rd!"), must_change_password=False)
    session.add_all([stock, admin])
    await session.flush()
    return co, cat, sub, cc, loc, dept, stock, admin


async def test_export_assets_xlsx_and_qr_png(client):
    async with SessionLocal() as session:
        co, cat, sub, cc, loc, dept, stock, it_admin = await _setup_company(session, "EXP1")
        rule = CodeRule(company_id=None, prefix_template="FA/{cost_center.code}/{category.code}/{subcategory.code}/CK_",
                         suffix_template="", start_number=1, pad_width=0)
        session.add(rule)
        await session.commit()
        assets = await procure_assets(session, {
            "company_id": co.id, "cost_center_id": cc.id, "category_id": cat.id, "subcategory_id": sub.id,
            "description": "Export Laptop", "purchase_date": date(2025, 12, 10), "initial_holder_id": stock.id,
        }, quantity=1, actor=it_admin)
        await session.commit()
        asset_id = assets[0].id
        asset_code = assets[0].asset_code

    resp = await client.post("/api/auth/login", json={"company_id": co.id, "emp_code": "ITA-EXP1", "password": "Passw0rd!"})
    headers = {"Authorization": f"Bearer {resp.json()['access_token']}"}

    export_resp = await client.get("/api/reports/export/assets", headers=headers)
    assert export_resp.status_code == 200
    assert export_resp.headers["content-type"].startswith("application/vnd.openxmlformats")

    wb = openpyxl.load_workbook(BytesIO(export_resp.content))
    ws = wb.active
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert header[0] == "Asset Code"
    codes_in_sheet = [row[0] for row in ws.iter_rows(min_row=2, values_only=True)]
    assert asset_code in codes_in_sheet

    qr_resp = await client.get(f"/api/assets/{asset_id}/qr.png", headers=headers)
    assert qr_resp.status_code == 200
    assert qr_resp.headers["content-type"] == "image/png"
    # PNG magic bytes -- confirms the response body is actually an image, not an
    # error payload masquerading under the right content-type header.
    assert qr_resp.content[:8] == b"\x89PNG\r\n\x1a\n"


async def test_export_assets_and_movements_scope_by_company():
    """An IT_TEAM (non-ADMIN) caller in Company A exporting the register or the
    movement log must never see Company B's asset codes -- same scoping guarantee
    as search_assets / dashboard_data (Task 19/20), just applied to the export
    endpoints."""
    from httpx import AsyncClient, ASGITransport
    from app.main import app
    from app.lifecycle.service import apply_event

    async with SessionLocal() as session:
        co_a, cat_a, sub_a, cc_a, loc, dept, stock_a, _admin_a = await _setup_company(session, "SC2A")
        co_b, cat_b, sub_b, cc_b, _loc_b, _dept_b, stock_b, admin_b = await _setup_company(session, "SC2B")
        # Company A's caller is IT_TEAM (not ADMIN) so scoped_company_ids actually
        # restricts it -- ADMIN is unrestricted by design, which would prove nothing.
        it_team_a = Holder(company_id=co_a.id, emp_code="ITT-SC2A", name="IT Team A", holder_type="EMPLOYEE",
                            location_id=loc.id, department_id=dept.id, role="IT_TEAM",
                            password_hash=hash_password("Passw0rd!"), must_change_password=False)
        emp_a = Holder(company_id=co_a.id, emp_code="EMP-SC2A", name="Employee A", holder_type="EMPLOYEE",
                        location_id=loc.id, department_id=dept.id, role="HOLDER")
        emp_b = Holder(company_id=co_b.id, emp_code="EMP-SC2B", name="Employee B", holder_type="EMPLOYEE",
                        location_id=loc.id, department_id=dept.id, role="HOLDER")
        session.add_all([it_team_a, emp_a, emp_b])
        rule = CodeRule(company_id=None, prefix_template="FA/{cost_center.code}/{category.code}/{subcategory.code}/CK_",
                         suffix_template="", start_number=1, pad_width=0)
        session.add(rule)
        await session.commit()

        [asset_a] = await procure_assets(session, {
            "company_id": co_a.id, "cost_center_id": cc_a.id, "category_id": cat_a.id, "subcategory_id": sub_a.id,
            "description": "Company A Laptop", "purchase_date": date(2025, 12, 10), "initial_holder_id": stock_a.id,
        }, quantity=1, actor=it_team_a)
        [asset_b] = await procure_assets(session, {
            "company_id": co_b.id, "cost_center_id": cc_b.id, "category_id": cat_b.id, "subcategory_id": sub_b.id,
            "description": "Company B Laptop", "purchase_date": date(2025, 12, 10), "initial_holder_id": stock_b.id,
        }, quantity=1, actor=admin_b)
        # A MOVED event dated "now" (apply_event's default) for each asset -- the
        # initial PROCURED event above is dated 2025-12-10 (its purchase_date), which
        # would fall outside the from/to window used below, so the assertions need a
        # same-day event to actually exercise the movements export's date filter.
        await apply_event(session, asset_a, "MOVED", to_holder_id=emp_a.id, actor=it_team_a)
        await apply_event(session, asset_b, "MOVED", to_holder_id=emp_b.id, actor=admin_b)
        await session.commit()
        asset_a_code, asset_b_code = asset_a.asset_code, asset_b.asset_code

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        resp = await client.post("/api/auth/login", json={"company_id": co_a.id, "emp_code": "ITT-SC2A", "password": "Passw0rd!"})
        headers = {"Authorization": f"Bearer {resp.json()['access_token']}"}

        assets_xlsx = (await client.get("/api/reports/export/assets", headers=headers)).content
        wb = openpyxl.load_workbook(BytesIO(assets_xlsx))
        codes = [row[0] for row in wb.active.iter_rows(min_row=2, values_only=True)]
        assert asset_a_code in codes
        assert asset_b_code not in codes

        from_date = (date.today() - timedelta(days=1)).isoformat()
        to_date = (date.today() + timedelta(days=1)).isoformat()
        movements_xlsx = (await client.get(
            f"/api/reports/export/movements?from_date={from_date}&to_date={to_date}", headers=headers,
        )).content
        wb2 = openpyxl.load_workbook(BytesIO(movements_xlsx))
        header2 = [c.value for c in next(wb2.active.iter_rows(min_row=1, max_row=1))]
        assert header2 == ["Asset Code", "Event", "Date", "From Holder", "To Holder", "Remarks"]
        move_rows = list(wb2.active.iter_rows(min_row=2, values_only=True))
        move_codes = [row[0] for row in move_rows]
        assert asset_a_code in move_codes
        assert asset_b_code not in move_codes

        # The "From Holder"/"To Holder" columns must hold the actual holder names
        # (joined in by the router), not the raw from_holder_id/to_holder_id -- a
        # regression back to raw ids would still pass the asset-code-only assertions
        # above, so this checks those two columns directly against the real names of
        # the holders the MOVED event above actually recorded.
        [moved_row_a] = [row for row in move_rows if row[0] == asset_a_code and row[1] == "MOVED"]
        from_holder_col, to_holder_col = moved_row_a[3], moved_row_a[4]
        assert from_holder_col == "IT Stock-SC2A"
        assert to_holder_col == "Employee A"
        assert not isinstance(from_holder_col, int)
        assert not isinstance(to_holder_col, int)


async def test_export_assets_pins_holder_to_only_their_own_held_asset():
    """A HOLDER calling GET /api/reports/export/assets must get exactly the asset(s)
    they currently hold -- never another employee's asset in the same company, and
    never anything from a different company. This is the same
    `holder_id = holder.id; allowed = None` branch `list_assets` (Task 19) uses, but
    it's a separately-maintained copy of that logic in the export endpoint, so it
    needs its own proof rather than relying on `list_assets`'s tests."""
    from httpx import AsyncClient, ASGITransport
    from app.main import app
    from app.lifecycle.service import apply_event

    async with SessionLocal() as session:
        co, cat, sub, cc, loc, dept, stock, admin = await _setup_company(session, "SC4A")
        co_other, cat_o, sub_o, cc_o, _loc_o, _dept_o, stock_o, admin_o = await _setup_company(session, "SC4B")

        holder_x = Holder(company_id=co.id, emp_code="HLDX-SC4", name="Holder X", holder_type="EMPLOYEE",
                           location_id=loc.id, department_id=dept.id, role="HOLDER",
                           password_hash=hash_password("Passw0rd!"), must_change_password=False)
        holder_y = Holder(company_id=co.id, emp_code="HLDY-SC4", name="Holder Y", holder_type="EMPLOYEE",
                           location_id=loc.id, department_id=dept.id, role="HOLDER",
                           password_hash=hash_password("Passw0rd!"), must_change_password=False)
        session.add_all([holder_x, holder_y])
        rule = CodeRule(company_id=None, prefix_template="FA/{cost_center.code}/{category.code}/{subcategory.code}/CK_",
                         suffix_template="", start_number=1, pad_width=0)
        session.add(rule)
        await session.commit()

        # Asset 1: same company, ends up held by holder_x (the caller below).
        [asset_x] = await procure_assets(session, {
            "company_id": co.id, "cost_center_id": cc.id, "category_id": cat.id, "subcategory_id": sub.id,
            "description": "Holder X's Laptop", "purchase_date": date(2025, 12, 10), "initial_holder_id": stock.id,
        }, quantity=1, actor=admin)
        await apply_event(session, asset_x, "MOVED", to_holder_id=holder_x.id, actor=admin)

        # Asset 2: same company, held by a *different* employee (holder_y) -- must not
        # appear in holder_x's export even though it's the same company.
        [asset_y] = await procure_assets(session, {
            "company_id": co.id, "cost_center_id": cc.id, "category_id": cat.id, "subcategory_id": sub.id,
            "description": "Holder Y's Laptop", "purchase_date": date(2025, 12, 10), "initial_holder_id": stock.id,
        }, quantity=1, actor=admin)
        await apply_event(session, asset_y, "MOVED", to_holder_id=holder_y.id, actor=admin)

        # Asset 3: a different company entirely.
        [asset_other_co] = await procure_assets(session, {
            "company_id": co_other.id, "cost_center_id": cc_o.id, "category_id": cat_o.id, "subcategory_id": sub_o.id,
            "description": "Other Company Laptop", "purchase_date": date(2025, 12, 10), "initial_holder_id": stock_o.id,
        }, quantity=1, actor=admin_o)

        await session.commit()
        code_x, code_y, code_other = asset_x.asset_code, asset_y.asset_code, asset_other_co.asset_code

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        resp = await client.post("/api/auth/login", json={"company_id": co.id, "emp_code": "HLDX-SC4", "password": "Passw0rd!"})
        headers = {"Authorization": f"Bearer {resp.json()['access_token']}"}

        export_resp = await client.get("/api/reports/export/assets", headers=headers)
        assert export_resp.status_code == 200
        wb = openpyxl.load_workbook(BytesIO(export_resp.content))
        codes = [row[0] for row in wb.active.iter_rows(min_row=2, values_only=True)]

        assert codes == [code_x]
        assert code_y not in codes
        assert code_other not in codes


async def test_qr_png_404_for_holder_who_does_not_hold_the_asset():
    """A HOLDER role calling the QR endpoint for an asset they do not currently hold
    must get 404, exactly like every other /api/assets/{id}/... route that goes
    through _get_scoped_asset (Task 16) -- a QR code is not an exception to that
    fail-closed scoping."""
    from httpx import AsyncClient, ASGITransport
    from app.main import app

    async with SessionLocal() as session:
        co, cat, sub, cc, loc, dept, stock, admin = await _setup_company(session, "SC3")
        other_holder = Holder(company_id=co.id, emp_code="EMP-SC3", name="Someone Else", holder_type="EMPLOYEE",
                               location_id=loc.id, department_id=dept.id, role="HOLDER",
                               password_hash=hash_password("Passw0rd!"), must_change_password=False)
        session.add(other_holder)
        rule = CodeRule(company_id=None, prefix_template="FA/{cost_center.code}/{category.code}/{subcategory.code}/CK_",
                         suffix_template="", start_number=1, pad_width=0)
        session.add(rule)
        await session.commit()
        [asset] = await procure_assets(session, {
            "company_id": co.id, "cost_center_id": cc.id, "category_id": cat.id, "subcategory_id": sub.id,
            "description": "Unheld Laptop", "purchase_date": date(2025, 12, 10), "initial_holder_id": stock.id,
        }, quantity=1, actor=admin)
        # asset's current holder is `stock`, not `other_holder`
        await session.commit()
        asset_id = asset.id

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        resp = await client.post("/api/auth/login", json={"company_id": co.id, "emp_code": "EMP-SC3", "password": "Passw0rd!"})
        headers = {"Authorization": f"Bearer {resp.json()['access_token']}"}

        qr_resp = await client.get(f"/api/assets/{asset_id}/qr.png", headers=headers)
        assert qr_resp.status_code == 404


def test_asset_qr_png_encodes_the_correct_asset_detail_url(monkeypatch):
    """Unit-checks that asset_qr_png hands the QR encoder exactly `{base_url}/assets/{id}`
    (not, say, a bare id or a different path) -- a QR decoder library isn't part of this
    project's dependencies, so this intercepts the call to qrcode.make() instead of round-
    tripping through actual pixel decoding."""
    from app.reports import export_service
    from app.core.config import settings

    captured = {}

    class FakeImg:
        def save(self, buf, format):  # noqa: A002 - matches PIL's Image.save signature
            buf.write(b"FAKE-PNG-BYTES")

    def fake_make(data):
        captured["data"] = data
        return FakeImg()

    monkeypatch.setattr(export_service.qrcode, "make", fake_make)

    result = export_service.asset_qr_png(42)

    assert captured["data"] == f"{settings.base_url}/assets/42"
    assert result == b"FAKE-PNG-BYTES"
