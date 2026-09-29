"""AM-06: the asset register export gains the full V1 field set (procurement,
PI Number, human-readable labels, Custom Field columns); a new, separate
field-change-audit export is added; the movement log export is confirmed
unchanged. See docs/ai/DECISIONS.md for why these are three separate
canonical datasets, never flattened into one another."""
from datetime import date
from io import BytesIO

import openpyxl
from app.assets.service import procure_assets
from app.core.db import SessionLocal
from app.core.security import hash_password
from app.holders.models import Holder
from app.masters.models import AssetCategory, AssetSubcategory, Brand, Company, CostCenter, CustomField, Department, Location, Vendor
from app.numbering.models import CodeRule


async def _setup(code="EXP06"):
    async with SessionLocal() as session:
        co = Company(code=f"{code}A", name=f"{code} Co A")
        co_b = Company(code=f"{code}B", name=f"{code} Co B")
        cat = AssetCategory(code=f"CAT-{code}", name="IT")
        vendor = Vendor(code=f"VND-{code}", name="Acme Traders")
        brand = Brand(code=f"BR1-{code}", name="Dell")
        brand2 = Brand(code=f"BR2-{code}", name="HP")
        session.add_all([co, co_b, cat, vendor, brand, brand2])
        await session.flush()
        sub = AssetSubcategory(category_id=cat.id, code="LAP", name="Laptop")
        cc = CostCenter(company_id=co.id, code="HO01", name="HO")
        loc = Location(code=f"HO-{code}", name="HO")
        dept = Department(name=f"IT-{code}")
        session.add_all([sub, cc, loc, dept])
        await session.flush()
        stock = Holder(company_id=co.id, emp_code=f"STK-{code}", name="IT Stock-HO", holder_type="IT_STOCK",
                        location_id=loc.id, department_id=dept.id, role="HOLDER")
        admin = Holder(company_id=co.id, emp_code=f"ADM-{code}", name="Admin", holder_type="EMPLOYEE",
                        location_id=loc.id, department_id=dept.id, role="ADMIN",
                        password_hash=hash_password("Passw0rd!"), must_change_password=False)
        # Company-scoped (never company_id=None/Global): a Global rule from one
        # _setup() call would otherwise be visible to (and, via
        # get_active_rule's "latest global rule wins" tiebreak, could even win
        # over) every OTHER test's company in this same file, cross-
        # contaminating which prefix/counter each test's assets actually use.
        rule = CodeRule(company_id=co.id, prefix_template=f"FA/{code}/", suffix_template="",
                         start_number=1, pad_width=0)
        session.add_all([stock, admin, rule])
        await session.commit()
        return {
            "co": co, "co_b_id": co_b.id, "cc": cc, "cat": cat, "sub": sub, "vendor": vendor,
            "brand": brand.id, "brand2": brand2.id,
            "stock": stock, "admin": admin,
        }


async def _headers(client, emp_code):
    resp = await client.post("/api/auth/login", json={"login_id": emp_code, "password": "Passw0rd!"})
    return {"Authorization": f"Bearer {resp.json()['access_token']}"}


def _header_row(ws) -> list:
    return [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]


class TestAssetRegisterExportColumns:
    async def test_export_includes_procurement_and_pi_number_with_human_readable_labels(self, client):
        ids = await _setup("COL1")
        async with SessionLocal() as session:
            [asset] = await procure_assets(session, {
                "company_id": ids["co"].id, "cost_center_id": ids["cc"].id, "category_id": ids["cat"].id,
                "subcategory_id": ids["sub"].id, "description": "Export Full Field Laptop",
                "purchase_date": date(2025, 6, 1), "initial_holder_id": ids["stock"].id,
                "vendor_id": ids["vendor"].id, "po_number": "PO-COL1", "po_date": date(2025, 5, 1),
                "invoice_number": "INV-COL1", "invoice_date": date(2025, 5, 2),
                "pi_number": "PI-COL1", "pi_date": date(2025, 5, 3),
                "purchase_cost": 60000, "tax_percent": 18,
                "brand_id": ids["brand"], "model": "Latitude", "serial_number": "SN-COL1",
                "warranty_upto": date(2027, 6, 1),
            }, quantity=1, actor=ids["admin"])
            await session.commit()
            asset_code = asset.asset_code

        headers = await _headers(client, ids["admin"].emp_code)
        resp = await client.get("/api/reports/export/assets", headers=headers)
        assert resp.status_code == 200
        wb = openpyxl.load_workbook(BytesIO(resp.content))
        ws = wb.active
        header = _header_row(ws)
        for col in ["Asset Code", "Company", "Cost Centre", "Category", "Subcategory", "Vendor",
                    "PO Number", "PO Date", "Invoice Number", "Invoice Date", "PI Number", "PI Date",
                    "Purchase Cost", "Tax Amount", "Total Cost", "Brand", "Model", "Serial Number", "Warranty Upto"]:
            assert col in header, f"missing column {col!r}"

        row_by_header = dict(zip(header, next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] == asset_code)))
        assert row_by_header["PI Number"] == "PI-COL1"
        assert row_by_header["Company"] == ids["co"].name  # human-readable, not a raw id
        assert row_by_header["Vendor"] == "Acme Traders"
        assert row_by_header["Current Holder"] == "IT Stock-HO"
        assert row_by_header["Holder Type"] == "IT_STOCK"

    async def test_export_includes_custom_field_columns_keyed_by_field_key(self, client):
        ids = await _setup("COL2")
        async with SessionLocal() as session:
            session.add(CustomField(field_key="col2_notes", label="Notes", field_type="text"))
            await session.commit()
        async with SessionLocal() as session:
            [asset] = await procure_assets(session, {
                "company_id": ids["co"].id, "cost_center_id": ids["cc"].id, "category_id": ids["cat"].id,
                "description": "Export UDF Laptop", "purchase_date": date(2025, 6, 1),
                "initial_holder_id": ids["stock"].id, "custom_fields": {"col2_notes": "hello export"},
            }, quantity=1, actor=ids["admin"])
            await session.commit()
            asset_code = asset.asset_code

        headers = await _headers(client, ids["admin"].emp_code)
        resp = await client.get("/api/reports/export/assets", headers=headers)
        wb = openpyxl.load_workbook(BytesIO(resp.content))
        ws = wb.active
        header = _header_row(ws)
        assert "Custom:col2_notes" in header
        row_by_header = dict(zip(header, next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] == asset_code)))
        assert row_by_header["Custom:col2_notes"] == "hello export"

    async def test_a_retired_custom_fields_stored_value_is_still_exported(self, client):
        """AM-06 §21: a value entered while a field was active/applicable
        must not disappear from the export just because the definition was
        later deactivated -- same never-silently-discard rule the AM-04
        Asset 360 tab already applies."""
        ids = await _setup("COL3")
        async with SessionLocal() as session:
            field = CustomField(field_key="col3_tag", label="Tag", field_type="text")
            session.add(field)
            await session.commit()
            field_id = field.id
        async with SessionLocal() as session:
            [asset] = await procure_assets(session, {
                "company_id": ids["co"].id, "cost_center_id": ids["cc"].id, "category_id": ids["cat"].id,
                "description": "Retired UDF Laptop", "purchase_date": date(2025, 6, 1),
                "initial_holder_id": ids["stock"].id, "custom_fields": {"col3_tag": "kept for export"},
            }, quantity=1, actor=ids["admin"])
            await session.commit()
            asset_code = asset.asset_code
        async with SessionLocal() as session:
            db_field = await session.get(CustomField, field_id)
            db_field.is_active = False
            await session.commit()

        headers = await _headers(client, ids["admin"].emp_code)
        resp = await client.get("/api/reports/export/assets", headers=headers)
        wb = openpyxl.load_workbook(BytesIO(resp.content))
        ws = wb.active
        header = _header_row(ws)
        assert "Custom:col3_tag" in header
        row_by_header = dict(zip(header, next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] == asset_code)))
        assert row_by_header["Custom:col3_tag"] == "kept for export"

    async def test_other_companys_custom_field_column_and_data_are_absent_from_a_scoped_export(self, client):
        ids = await _setup("COL4")
        async with SessionLocal() as session:
            session.add(CustomField(field_key="col4_other", label="Other", field_type="text", company_id=ids["co_b_id"]))
            it_team = Holder(company_id=ids["co"].id, emp_code="ITT-COL4", name="IT Team", holder_type="EMPLOYEE",
                              location_id=ids["stock"].location_id, department_id=ids["stock"].department_id,
                              role="IT_TEAM", password_hash=hash_password("Passw0rd!"), must_change_password=False)
            session.add(it_team)
            await session.commit()
        async with SessionLocal() as session:
            [asset] = await procure_assets(session, {
                "company_id": ids["co"].id, "cost_center_id": ids["cc"].id, "category_id": ids["cat"].id,
                "description": "Company A Only Laptop", "purchase_date": date(2025, 6, 1),
                "initial_holder_id": ids["stock"].id,
            }, quantity=1, actor=ids["admin"])
            await session.commit()

        headers = await _headers(client, "ITT-COL4")
        resp = await client.get("/api/reports/export/assets", headers=headers)
        wb = openpyxl.load_workbook(BytesIO(resp.content))
        header = _header_row(wb.active)
        assert "Custom:col4_other" not in header


class TestFieldChangeAuditExport:
    async def test_export_contains_asset_code_actor_old_new_values_chronologically(self, client):
        ids = await _setup("AUD1")
        async with SessionLocal() as session:
            [asset] = await procure_assets(session, {
                "company_id": ids["co"].id, "cost_center_id": ids["cc"].id, "category_id": ids["cat"].id,
                "description": "Audit Export Laptop", "purchase_date": date(2025, 6, 1),
                "initial_holder_id": ids["stock"].id, "brand_id": ids["brand"],
            }, quantity=1, actor=ids["admin"])
            await session.commit()
            asset_id, asset_code = asset.id, asset.asset_code

        # AssetUpdateIn is a full-replace PUT (docs/ai/DECISIONS.md): a field
        # omitted from the body resets to its schema default, not "left
        # unchanged" -- every editable field must be resubmitted each time.
        headers = await _headers(client, ids["admin"].emp_code)
        full_body = {
            "legacy_asset_code": None, "brand_id": ids["brand"], "model": None, "serial_number": None,
            "description": "Audit Export Laptop", "vendor_id": None, "po_number": None, "po_date": None,
            "invoice_number": None, "invoice_date": None, "pi_number": None, "pi_date": None,
            # procure_assets defaults an unset cost/tax to 0 (Decimal), not None --
            # these must match that, or the very first no-op field lands as a
            # spurious 0.00 -> None "change".
            "purchase_cost": 0, "tax_percent": 0, "warranty_years": None, "custom_fields": None,
        }
        await client.put(f"/api/assets/{asset_id}", json={**full_body, "brand_id": ids["brand2"]}, headers=headers)
        await client.put(f"/api/assets/{asset_id}", json={**full_body, "brand_id": ids["brand2"], "description": "Audit Export Laptop v2"}, headers=headers)

        resp = await client.get("/api/reports/export/field-changes", headers=headers)
        assert resp.status_code == 200
        wb = openpyxl.load_workbook(BytesIO(resp.content))
        ws = wb.active
        assert _header_row(ws) == ["Asset Code", "Field", "Old Value", "New Value", "Actor", "Changed At", "Request ID", "Reason"]
        rows = [dict(zip(_header_row(ws), r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[0] == asset_code]
        assert len(rows) == 2  # one brand change, one description change (in that PUT order)
        assert rows[0]["Field"] == "brand_id"
        assert rows[0]["Old Value"] == f"Dell (#{ids['brand']})" and rows[0]["New Value"] == f"HP (#{ids['brand2']})"
        assert rows[0]["Actor"] == "Admin"
        # AM-17 DEF-02: an ordinary edit's rows carry no Reason (that column
        # is what distinguishes them from a controlled correction).
        assert rows[0]["Reason"] is None
        assert rows[1]["Field"] == "description"
        # chronological: row 0's Changed At <= row 1's Changed At
        assert rows[0]["Changed At"] <= rows[1]["Changed At"]

    async def test_field_change_export_includes_correction_reason(self, client):
        """AM-17 DEF-02: a controlled correction's Reason (the sole
        discriminator between a correction row and an ordinary edit row) was
        missing from this export entirely."""
        ids = await _setup("AUD-REASON")
        async with SessionLocal() as session:
            [asset] = await procure_assets(session, {
                "company_id": ids["co"].id, "cost_center_id": ids["cc"].id, "category_id": ids["cat"].id,
                "description": "Reason Export Laptop", "purchase_date": date(2025, 6, 1),
                "initial_holder_id": ids["stock"].id,
            }, quantity=1, actor=ids["admin"])
            await session.commit()
            asset_id, asset_code = asset.id, asset.asset_code

        headers = await _headers(client, ids["admin"].emp_code)
        resp = await client.post(
            f"/api/assets/{asset_id}/corrections",
            json={"purchase_date": "2025-05-15", "reason": "Invoice date was mis-keyed at entry"},
            headers=headers,
        )
        assert resp.status_code == 200

        resp = await client.get("/api/reports/export/field-changes", headers=headers)
        wb = openpyxl.load_workbook(BytesIO(resp.content))
        ws = wb.active
        rows = [dict(zip(_header_row(ws), r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[0] == asset_code]
        assert len(rows) == 1
        assert rows[0]["Field"] == "purchase_date"
        assert rows[0]["Reason"] == "Invoice date was mis-keyed at entry"

    async def test_field_change_export_is_company_scoped(self, client):
        # Company B's caller is IT_TEAM, not ADMIN -- scoped_company_ids
        # returns None (unrestricted) for ADMIN by design, which would prove
        # nothing about scoping (same trap test_export.py's own comment
        # documents for the asset/movement exports).
        ids_a = await _setup("AUD2A")
        ids_b = await _setup("AUD2B")
        async with SessionLocal() as session:
            it_team_b = Holder(company_id=ids_b["co"].id, emp_code="ITT-AUD2B", name="IT Team B",
                                holder_type="EMPLOYEE", location_id=ids_b["stock"].location_id,
                                department_id=ids_b["stock"].department_id, role="IT_TEAM",
                                password_hash=hash_password("Passw0rd!"), must_change_password=False)
            session.add(it_team_b)
            await session.commit()
        async with SessionLocal() as session:
            [asset_a] = await procure_assets(session, {
                "company_id": ids_a["co"].id, "cost_center_id": ids_a["cc"].id, "category_id": ids_a["cat"].id,
                "description": "A's Laptop", "purchase_date": date(2025, 6, 1),
                "initial_holder_id": ids_a["stock"].id,
            }, quantity=1, actor=ids_a["admin"])
            await session.commit()
            asset_a_id, asset_a_code = asset_a.id, asset_a.asset_code

        headers_a = await _headers(client, ids_a["admin"].emp_code)
        await client.put(f"/api/assets/{asset_a_id}", json={"description": "changed by A"}, headers=headers_a)

        headers_b = await _headers(client, "ITT-AUD2B")
        resp = await client.get("/api/reports/export/field-changes", headers=headers_b)
        wb = openpyxl.load_workbook(BytesIO(resp.content))
        codes = [r[0] for r in wb.active.iter_rows(min_row=2, values_only=True)]
        assert asset_a_code not in codes

    async def test_holder_role_cannot_export_field_changes(self, client):
        ids = await _setup("AUD3")
        async with SessionLocal() as session:
            holder_role = Holder(company_id=ids["co"].id, emp_code="HLD-AUD3", name="Just A Holder",
                                  holder_type="EMPLOYEE", location_id=ids["stock"].location_id,
                                  department_id=ids["stock"].department_id, role="HOLDER",
                                  password_hash=hash_password("Passw0rd!"), must_change_password=False)
            session.add(holder_role)
            await session.commit()
        headers = await _headers(client, "HLD-AUD3")
        resp = await client.get("/api/reports/export/field-changes", headers=headers)
        assert resp.status_code == 403


class TestMovementExportUnchanged:
    async def test_movement_export_still_uses_point_in_time_holder_snapshots(self, client):
        """AM-06 makes no change to the movement log -- reconfirms the AM-01
        snapshot behavior the export already relied on still holds."""
        from datetime import timedelta
        from app.lifecycle.service import apply_event

        ids = await _setup("MOV1")
        async with SessionLocal() as session:
            employee = Holder(company_id=ids["co"].id, emp_code="EMP-MOV1", name="Original Name",
                               holder_type="EMPLOYEE", location_id=ids["stock"].location_id,
                               department_id=ids["stock"].department_id, role="HOLDER")
            session.add(employee)
            await session.commit()
        async with SessionLocal() as session:
            [asset] = await procure_assets(session, {
                "company_id": ids["co"].id, "cost_center_id": ids["cc"].id, "category_id": ids["cat"].id,
                "description": "Movement Snapshot Laptop", "purchase_date": date(2025, 6, 1),
                "initial_holder_id": ids["stock"].id,
            }, quantity=1, actor=ids["admin"])
            await apply_event(session, asset, "MOVED", to_holder_id=employee.id, actor=ids["admin"])
            await session.commit()
            asset_code = asset.asset_code

        # Renaming the holder AFTER the move must not retroactively change the
        # exported "To Holder" value (the AM-01 point-in-time snapshot rule).
        async with SessionLocal() as session:
            db_employee = await session.get(Holder, employee.id)
            db_employee.name = "Renamed Later"
            await session.commit()

        headers = await _headers(client, ids["admin"].emp_code)
        today = date.today()
        resp = await client.get(
            f"/api/reports/export/movements?from_date={(today - timedelta(days=1)).isoformat()}"
            f"&to_date={(today + timedelta(days=1)).isoformat()}",
            headers=headers,
        )
        wb = openpyxl.load_workbook(BytesIO(resp.content))
        moved_row = next(r for r in wb.active.iter_rows(min_row=2, values_only=True) if r[0] == asset_code and r[1] == "MOVED")
        assert moved_row[4] == "Original Name"
