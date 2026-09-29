"""AM-23: Import's new Edit mode -- bulk-corrects existing assets by Asset
Code. A blank cell means "leave this field exactly as it is", never
"clear it" -- see app.imports.asset_import_service's module docstring and
`_validate_edit_rows`."""
import io
from datetime import date
import openpyxl
from app.assets.models import Asset, AssetFieldChange
from app.assets.service import procure_assets
from app.core.db import SessionLocal
from app.core.security import hash_password
from app.asset_users.models import AssetUser
from app.masters.models import AssetCategory, AssetSubcategory, Brand, Company, CostCenter, CustomField, Department, Location, Vendor
from app.numbering.models import CodeRule
from sqlalchemy import select

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

EDIT_HEADER = [
    "Asset Code", "Legacy Asset Code", "Brand Code", "Model", "Serial Number", "Barcode",
    "Description", "Vendor Code", "PO Number", "PO Date", "Invoice Number", "Invoice Date",
    "Invoice Amount", "PI Number", "PI Date", "Purchase Cost", "Tax %", "Warranty Years",
]


def _xlsx(rows: list[dict], extra_headers: list[str] | None = None) -> bytes:
    header = EDIT_HEADER + (extra_headers or [])
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(header)
    for row in rows:
        ws.append([row.get(h) for h in header])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def _post(client, path, content, headers, mode="edit"):
    return await client.post(f"{path}?mode={mode}", files={"file": ("a.xlsx", io.BytesIO(content), XLSX)}, headers=headers)


async def _setup(suffix: str):
    async with SessionLocal() as session:
        co = Company(code=f"AM23E-{suffix}", name=f"AM23 Edit Co {suffix}")
        co_b = Company(code=f"AM23E-{suffix}-B", name=f"AM23 Edit Co {suffix} B")
        cat = AssetCategory(code=f"AM23E-{suffix}", name="IT", asset_domain="IT")
        vendor = Vendor(code=f"VND-{suffix}", name="Vendor One")
        vendor2 = Vendor(code=f"VND2-{suffix}", name="Vendor Two")
        brand = Brand(code=f"BR-{suffix}", name="OldBrand")
        brand2 = Brand(code=f"BR2-{suffix}", name="NewBrand")
        session.add_all([co, co_b, cat, vendor, vendor2, brand, brand2])
        await session.flush()
        sub = AssetSubcategory(category_id=cat.id, code="LAP", name="Laptop")
        cc = CostCenter(company_id=co.id, code="HO", name="Head Office")
        loc = Location(company_id=co.id, code=f"AM23E-{suffix}", name="HO")
        dept = Department(name=f"AM23E-{suffix}")
        session.add_all([sub, cc, loc, dept])
        await session.flush()
        stock = AssetUser(company_id=co.id, code=f"STK-{suffix}", name="IT Stock", asset_user_type="STOCK_POINT",
                        location_id=loc.id, department_id=dept.id, role="SELF_SERVICE")
        # Bulk asset Import is Primary-Owner-only now, not merely ADMIN-only.
        admin = AssetUser(company_id=co.id, code=f"ADM-{suffix}", name="Admin", asset_user_type="EMPLOYEE",
                        location_id=loc.id, department_id=dept.id, role="ADMIN", is_primary_owner=True,
                        login_enabled=True, password_hash=hash_password("Passw0rd!"), must_change_password=False)
        ita = AssetUser(company_id=co.id, code=f"ITA-{suffix}", name="IT Team", asset_user_type="EMPLOYEE",
                     location_id=loc.id, department_id=dept.id, role="OPERATOR",
                     login_enabled=True, password_hash=hash_password("Passw0rd!"), must_change_password=False)
        rule = CodeRule(company_id=co.id, prefix_template=f"AM23E/{suffix}/", suffix_template="",
                         start_number=1, pad_width=0)
        session.add_all([stock, admin, ita, rule])
        await session.commit()

        [asset] = await procure_assets(session, {
            "company_id": co.id, "cost_center_id": cc.id, "category_id": cat.id, "subcategory_id": sub.id,
            "description": "Original Description", "purchase_date": date(2025, 6, 1),
            "serial_number": f"SN-{suffix}-ORIG", "initial_asset_user_id": stock.id,
            "brand_id": brand.id, "model": "OldModel", "vendor_id": vendor.id,
            "purchase_cost": 1000, "tax_percent": 10, "warranty_years": 2,
        }, quantity=1, actor=admin)
        await session.commit()
        return {
            "co": co, "co_b": co_b, "admin": f"ADM-{suffix}", "ita": f"ITA-{suffix}",
            "asset_id": asset.id, "asset_code": asset.asset_code,
            "vendor": vendor.code, "vendor2": vendor2.code,
            "brand": brand.code, "brand_id": brand.id, "brand2": brand2.code, "brand2_id": brand2.id,
        }


async def _login(client, co_id, login_id):
    resp = await client.post("/api/auth/login", json={"company_id": co_id, "login_id": login_id, "password": "Passw0rd!"})
    assert resp.status_code == 200, resp.text
    return {"Authorization": f"Bearer {resp.json()['access_token']}"}


async def _get_asset(asset_id: int) -> Asset:
    async with SessionLocal() as session:
        return await session.get(Asset, asset_id)


async def test_edit_template_has_asset_code_and_no_creation_only_columns(client):
    ids = await _setup("TPL1")
    headers = await _login(client, ids["co"].id, ids["admin"])
    resp = await client.get("/api/imports/assets/template?mode=edit", headers=headers)
    assert resp.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    header = [c.value for c in next(wb.active.iter_rows(min_row=1, max_row=1))]
    assert "Asset Code" in header
    for creation_only in ("Company Code", "Cost Centre Code", "Category Code", "Subcategory Code",
                           "Initial AssetUser Code", "Quantity"):
        assert creation_only not in header


async def test_blank_cells_leave_every_other_field_untouched(client):
    """The core Edit-mode guarantee: a row that only fills in PI Number
    must not blank out Brand/Model/Description/anything else."""
    ids = await _setup("BLANK1")
    headers = await _login(client, ids["co"].id, ids["admin"])
    row = {"Asset Code": ids["asset_code"], "PI Number": "PI-999", "PI Date": "2026-01-05"}

    resp = await _post(client, "/api/imports/assets/commit", _xlsx([row]), headers)
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["updated"] == 1
    assert body["errors"] == []

    asset = await _get_asset(ids["asset_id"])
    assert asset.pi_number == "PI-999"
    assert asset.pi_date == date(2026, 1, 5)
    assert asset.brand_id == ids["brand_id"]
    assert asset.model == "OldModel"
    assert asset.description == "Original Description"
    assert float(asset.purchase_cost) == 1000.0
    assert float(asset.tax_percent) == 10.0


async def test_provided_fields_are_updated_together(client):
    ids = await _setup("MULTI1")
    headers = await _login(client, ids["co"].id, ids["admin"])
    row = {
        "Asset Code": ids["asset_code"], "Brand Code": ids["brand2"], "Model": "NewModel",
        "Description": "Updated Description", "Vendor Code": ids["vendor2"],
        "Purchase Cost": 2000, "Tax %": 20,
    }
    resp = await _post(client, "/api/imports/assets/commit", _xlsx([row]), headers)
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["updated"] == 1

    asset = await _get_asset(ids["asset_id"])
    assert asset.brand_id == ids["brand2_id"]
    assert asset.model == "NewModel"
    assert asset.description == "Updated Description"
    assert float(asset.purchase_cost) == 2000.0
    assert float(asset.tax_percent) == 20.0
    # Tax recomputed from the new effective cost/percent.
    assert float(asset.tax_amount) == 400.0
    assert float(asset.total_cost) == 2400.0


async def test_warranty_years_only_recomputes_warranty_upto_when_provided(client):
    ids = await _setup("WARR1")
    headers = await _login(client, ids["co"].id, ids["admin"])
    before = await _get_asset(ids["asset_id"])
    assert before.warranty_upto == date(2027, 5, 31)  # 2025-06-01 + 2y - 1d

    # Row 1: no Warranty Years mentioned -- untouched.
    resp = await _post(client, "/api/imports/assets/commit", _xlsx([{"Asset Code": ids["asset_code"], "Brand Code": ids["brand2"]}]), headers)
    assert resp.status_code == 200
    mid = await _get_asset(ids["asset_id"])
    assert mid.warranty_years == 2
    assert mid.warranty_upto == date(2027, 5, 31)

    # Row 2: Warranty Years explicitly changed -- recomputed from the
    # asset's own (unchanged) purchase_date.
    resp = await _post(client, "/api/imports/assets/commit", _xlsx([{"Asset Code": ids["asset_code"], "Warranty Years": 5}]), headers)
    assert resp.status_code == 200
    after = await _get_asset(ids["asset_id"])
    assert after.warranty_years == 5
    assert after.warranty_upto == date(2030, 5, 31)


async def test_serial_number_uniqueness_is_enforced_and_excludes_self(client):
    ids = await _setup("SER1")
    headers = await _login(client, ids["co"].id, ids["admin"])

    # Re-submitting the asset's own current serial number is a no-op, not a conflict.
    resp = await _post(client, "/api/imports/assets/commit",
                        _xlsx([{"Asset Code": ids["asset_code"], "Serial Number": f"SN-SER1-ORIG"}]), headers)
    assert resp.status_code == 200
    assert resp.json()["updated"] == 1

    # A second asset already holds "TAKEN" -- editing onto it is rejected as a row error.
    async with SessionLocal() as session:
        co = await session.get(Company, ids["co"].id)
        stock = (await session.execute(select(AssetUser).where(AssetUser.company_id == co.id, AssetUser.asset_user_type == "STOCK_POINT"))).scalars().first()
        cc = (await session.execute(select(CostCenter).where(CostCenter.company_id == co.id))).scalars().first()
        cat = (await session.execute(select(AssetCategory).where(AssetCategory.code == f"AM23E-SER1"))).scalars().first()
        admin = (await session.execute(select(AssetUser).where(AssetUser.company_id == co.id, AssetUser.code == "ADM-SER1"))).scalars().first()
        await procure_assets(session, {
            "company_id": co.id, "cost_center_id": cc.id, "category_id": cat.id, "subcategory_id": None,
            "description": "Other asset", "purchase_date": date(2025, 1, 1),
            "serial_number": "TAKEN-SER1", "initial_asset_user_id": stock.id,
        }, quantity=1, actor=admin)
        await session.commit()

    resp = await _post(client, "/api/imports/assets/commit",
                        _xlsx([{"Asset Code": ids["asset_code"], "Serial Number": "TAKEN-SER1"}]), headers)
    assert resp.status_code == 200
    body = resp.json()
    assert body["updated"] == 0
    assert "serial number" in body["errors"][0]["message"].lower()
    asset = await _get_asset(ids["asset_id"])
    assert asset.serial_number == "SN-SER1-ORIG"


async def test_unknown_asset_code_is_a_row_error(client):
    ids = await _setup("UNK1")
    headers = await _login(client, ids["co"].id, ids["admin"])
    resp = await _post(client, "/api/imports/assets/commit", _xlsx([{"Asset Code": "NOPE-DOES-NOT-EXIST", "Brand Code": ids["brand2"]}]), headers)
    assert resp.status_code == 200
    body = resp.json()
    assert body["updated"] == 0
    assert "unknown Asset Code" in body["errors"][0]["message"]


async def test_a_row_with_no_fields_filled_in_besides_asset_code_is_a_row_error(client):
    ids = await _setup("EMPTY1")
    headers = await _login(client, ids["co"].id, ids["admin"])
    resp = await _post(client, "/api/imports/assets/commit", _xlsx([{"Asset Code": ids["asset_code"]}]), headers)
    assert resp.status_code == 200
    body = resp.json()
    assert body["updated"] == 0
    assert "nothing to update" in body["errors"][0]["message"]


async def test_operator_cannot_edit_via_import_at_all_even_within_its_own_company(client):
    """Bulk asset Import is Primary-Owner-only (rebuild rule) -- OPERATOR
    gets 403 unconditionally, not just when the target asset is out of scope."""
    ids = await _setup("SCOPE1")
    headers = await _login(client, ids["co"].id, ids["ita"])
    resp = await _post(client, "/api/imports/assets/commit", _xlsx([{"Asset Code": ids["asset_code"], "Brand Code": ids["brand2"]}]), headers)
    assert resp.status_code == 403

    # A second asset that belongs to a DIFFERENT company -- refused too, still 403.
    async with SessionLocal() as session:
        co_b = await session.get(Company, ids["co_b"].id)
        loc = Location(company_id=co_b.id, code="SCOPE1-B", name="B HO")
        dept = Department(name="SCOPE1-B")
        session.add_all([loc, dept])
        await session.flush()
        cc = CostCenter(company_id=co_b.id, code="HO", name="HO")
        cat = AssetCategory(code="SCOPE1-B", name="IT", asset_domain="IT")
        session.add_all([cc, cat])
        await session.flush()
        stock_b = AssetUser(company_id=co_b.id, code="STKB-SCOPE1", name="B Stock", asset_user_type="STOCK_POINT",
                          location_id=loc.id, department_id=dept.id, role="SELF_SERVICE")
        admin_b = AssetUser(company_id=co_b.id, code="ADMB-SCOPE1", name="B Admin", asset_user_type="EMPLOYEE",
                          location_id=loc.id, department_id=dept.id, role="ADMIN",
                          login_enabled=True, password_hash=hash_password("Passw0rd!"), must_change_password=False)
        rule_b = CodeRule(company_id=co_b.id, prefix_template="SCOPE1B/", suffix_template="", start_number=1, pad_width=0)
        session.add_all([stock_b, admin_b, rule_b])
        await session.commit()
        [asset_b] = await procure_assets(session, {
            "company_id": co_b.id, "cost_center_id": cc.id, "category_id": cat.id, "subcategory_id": None,
            "description": "Company B asset", "purchase_date": date(2025, 1, 1),
            "serial_number": "SN-SCOPE1-B", "initial_asset_user_id": stock_b.id,
        }, quantity=1, actor=admin_b)
        await session.commit()
        asset_b_code = asset_b.asset_code

    resp = await _post(client, "/api/imports/assets/commit", _xlsx([{"Asset Code": asset_b_code, "Brand Code": ids["brand2"]}]), headers)
    assert resp.status_code == 403


async def test_custom_field_values_are_merged_not_replaced(client):
    ids = await _setup("UDF1")
    async with SessionLocal() as session:
        session.add(CustomField(field_key="udf1_notes", label="Notes", field_type="text"))
        session.add(CustomField(field_key="udf1_tag", label="Tag", field_type="text"))
        await session.commit()
        asset = await session.get(Asset, ids["asset_id"])
        asset.custom_fields = {"udf1_notes": "existing note"}
        await session.commit()

    headers = await _login(client, ids["co"].id, ids["admin"])
    row = {"Asset Code": ids["asset_code"], "Custom:udf1_tag": "TAG-1"}
    resp = await _post(client, "/api/imports/assets/commit", _xlsx([row], extra_headers=["Custom:udf1_tag"]), headers)
    assert resp.status_code == 200, resp.text
    assert resp.json()["updated"] == 1

    asset = await _get_asset(ids["asset_id"])
    assert asset.custom_fields["udf1_notes"] == "existing note"
    assert asset.custom_fields["udf1_tag"] == "TAG-1"


async def test_edit_writes_an_audit_row_only_for_genuinely_changed_fields(client):
    ids = await _setup("AUDIT1")
    headers = await _login(client, ids["co"].id, ids["admin"])
    row = {"Asset Code": ids["asset_code"], "Brand Code": ids["brand2"], "Model": "OldModel"}  # Model unchanged
    resp = await _post(client, "/api/imports/assets/commit", _xlsx([row]), headers)
    assert resp.status_code == 200
    assert resp.json()["updated"] == 1

    async with SessionLocal() as session:
        changes = (await session.execute(
            select(AssetFieldChange).where(AssetFieldChange.asset_id == ids["asset_id"])
        )).scalars().all()
    fields_changed = {c.field_name for c in changes}
    assert "brand_id" in fields_changed
    assert "model" not in fields_changed


async def test_preview_never_writes_to_the_database(client):
    ids = await _setup("PREV1")
    headers = await _login(client, ids["co"].id, ids["admin"])
    resp = await _post(client, "/api/imports/assets/preview", _xlsx([{"Asset Code": ids["asset_code"], "Brand Code": ids["brand2"]}]), headers)
    assert resp.status_code == 200
    body = resp.json()
    assert len(body["valid_rows"]) == 1
    assert body["valid_rows"][0]["asset_code"] == ids["asset_code"]
    assert "brand_id" in body["valid_rows"][0]["fields_changed"]

    asset = await _get_asset(ids["asset_id"])
    assert asset.brand_id == ids["brand_id"]  # untouched
