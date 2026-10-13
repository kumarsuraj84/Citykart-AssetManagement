"""S7 fix round (P2-10): the asset Excel import (add mode) works like Add Assets -- an optional "Item" column
(the Item decides Category, Sub-category and work area; Category / Subcategory Code may then stay empty),
Vendor Code optional, and the column "Received Into Code" (the old "Initial AssetUser Code" still read).
The same import row engine creates the assets."""
import io

import openpyxl
from sqlalchemy import select

from app.assets.models import Asset
from app.core.db import SessionLocal
from app.imports.asset_import_service import ADD_TEMPLATE_COLUMNS, _validate_add_rows
from app.items.models import Item
from tests.purchase_orders.s4_world import build_s4_world, login

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
NEW = ["Company Code", "Cost Centre Code", "Item", "Category Code", "Subcategory Code", "Description", "Vendor Code",
       "Serial Number", "Received Into Code", "Quantity", "Brand Code", "Warranty Years", "Legacy Asset Code"]


def xlsx(rows: list[dict], header: list[str] = NEW) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(header)
    for row in rows:
        ws.append([row.get(h) for h in header])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def row(**over) -> dict:
    base = {"Company Code": "S0A", "Cost Centre Code": "HO", "Serial Number": "N/A", "Received Into Code": "S0STK"}
    base.update(over)
    return base


async def post(client, h, content: bytes, step: str = "commit"):
    return await client.post(f"/api/imports/assets/{step}", headers=h, files={"file": ("a.xlsx", io.BytesIO(content), XLSX)})


async def assets() -> list[Asset]:
    async with SessionLocal() as s:
        return list((await s.execute(select(Asset).order_by(Asset.id))).scalars().all())


async def test_the_template_has_the_item_column_and_the_new_heading(client):
    w = await build_s4_world()
    resp = await client.get("/api/imports/assets/template", headers=await login(client, w, "owner"))
    header = [c.value for c in next(openpyxl.load_workbook(io.BytesIO(resp.content)).active.iter_rows(max_row=1))]
    assert header == ADD_TEMPLATE_COLUMNS
    assert "Item" in header and "Received Into Code" in header and "Initial AssetUser Code" not in header


async def test_an_item_row_needs_no_category_description_or_vendor(client):
    w = await build_s4_world()
    h = await login(client, w, "owner")
    resp = await post(client, h, xlsx([row(Item=" office CHAIR ", Quantity=3)]))
    assert resp.status_code == 200, resp.text
    assert (resp.json()["imported"], resp.json()["errors"]) == (3, [])
    made = await assets()
    assert len(made) == 3
    for a in made:
        assert (a.item_id, a.category_id, a.subcategory_id, a.asset_domain) == (w["item_chair"], w["nonit"], w["chair"], "NON_IT")
        assert a.description == "Office Chair" and a.vendor_id is None and a.current_asset_user_id == w["stock"]


async def test_the_preview_names_the_item_of_each_row(client):
    w = await build_s4_world()
    h = await login(client, w, "owner")
    body = (await post(client, h, xlsx([row(Item="Gang Box"), row(Item="Gang Box", Description="Box in the lobby")]), "preview")).json()
    assert [r["item"] for r in body["valid_rows"]] == ["Gang Box", "Gang Box"]
    assert [r["category"] for r in body["valid_rows"]] == ["Computers", "Computers"]
    assert [r["description"] for r in body["valid_rows"]] == ["Gang Box", "Box in the lobby"]
    assert await assets() == []                                                          # a preview writes nothing


async def test_an_unknown_or_retired_item_is_a_plain_row_error(client):
    w = await build_s4_world()
    h = await login(client, w, "owner")
    async with SessionLocal() as s:
        (await s.get(Item, w["item_ap"])).is_active = False
        await s.commit()
    body = (await post(client, h, xlsx([row(Item="Flux Capacitor"), row(Item="Access Point"), row(Item="Gang Box")]), "preview")).json()
    assert [e["row"] for e in body["errors"]] == [2, 3]
    assert body["errors"][0]["message"] == ("There is no active Item named 'Flux Capacitor'. Check the name, or add the Item "
                                            "under Settings > Items first.")
    assert body["errors"][0]["field"] == "Item" and "Access Point" in body["errors"][1]["message"]
    assert [r["row"] for r in body["valid_rows"]] == [4]


async def test_a_category_that_disagrees_with_the_item_is_refused(client):
    w = await build_s4_world()
    h = await login(client, w, "owner")
    body = (await post(client, h, xlsx([
        row(Item="Office Chair", **{"Category Code": "S0IT"}),
        row(Item="Office Chair", **{"Category Code": "s0ni", "Subcategory Code": "LAP"}),
        row(Item="Office Chair", **{"Category Code": "S0NI", "Subcategory Code": "chr"}),
    ]), "preview")).json()
    assert body["errors"][0]["message"] == ("The Item 'Office Chair' belongs to the category Furniture (S0NI), not 'S0IT'. "
                                            "Leave Category Code empty or correct it.")
    assert "belongs to the sub-category Chair (CHR), not 'LAP'" in body["errors"][1]["message"]
    assert [r["row"] for r in body["valid_rows"]] == [4]                                 # matching codes are fine


async def test_without_an_item_the_category_columns_work_as_before_and_vendor_is_optional(client):
    w = await build_s4_world()
    h = await login(client, w, "owner")
    resp = await post(client, h, xlsx([row(**{"Category Code": "S0IT", "Subcategory Code": "LAP", "Description": "Old laptop"})]))
    assert resp.json()["imported"] == 1
    [a] = await assets()
    assert (a.item_id, a.category_id, a.vendor_id, a.description) == (None, w["it"], None, "Old laptop")
    missing = (await post(client, h, xlsx([row(Description="Something")]), "preview")).json()
    assert missing["errors"][0]["message"] == "Give the Item (or the Category Code and Subcategory Code)"


async def test_a_file_from_the_old_template_still_imports(client):
    w = await build_s4_world()
    h = await login(client, w, "owner")
    old = ["Company Code", "Cost Centre Code", "Category Code", "Subcategory Code", "Description", "Vendor Code",
           "Serial Number", "Initial AssetUser Code"]
    content = xlsx([{"Company Code": "S0A", "Cost Centre Code": "HO", "Category Code": "S0IT", "Subcategory Code": "LAP",
                     "Description": "Laptop", "Vendor Code": "S0V", "Serial Number": "OLD-1", "Initial AssetUser Code": "S0STK"},
                    {"Company Code": "S0A", "Cost Centre Code": "HO", "Category Code": "S0IT", "Subcategory Code": "LAP",
                     "Description": "Laptop", "Vendor Code": "S0V", "Serial Number": "OLD-2", "Initial AssetUser Code": "NOPE"}], old)
    body = (await post(client, h, content)).json()
    assert body["imported"] == 1
    assert body["errors"] == [{"row": 3, "field": "Initial AssetUser Code", "message": "unknown Initial AssetUser Code 'NOPE'"}]
    new_header = (await post(client, h, xlsx([row(Item="Gang Box", **{"Received Into Code": "NOPE"})]), "preview")).json()
    assert new_header["errors"][0] == {"row": 2, "field": "Received Into Code", "message": "unknown Received Into Code 'NOPE'"}


async def test_a_file_with_neither_items_nor_categories_is_refused_as_a_whole(client):
    w = await build_s4_world()
    h = await login(client, w, "owner")
    resp = await post(client, h, xlsx([row()], ["Company Code", "Cost Centre Code", "Serial Number", "Received Into Code"]), "preview")
    assert resp.status_code == 422
    assert "Category Code" in resp.json()["detail"] and "Description" in resp.json()["detail"]


async def test_the_items_defaults_apply_like_add_assets(client):
    w = await build_s4_world()
    h = await login(client, w, "owner")
    async with SessionLocal() as s:
        chair = await s.get(Item, w["item_chair"])
        chair.default_brand_id, chair.default_warranty_years = w["brand"], 2
        await s.commit()
    assert (await post(client, h, xlsx([row(Item="Office Chair"), row(Item="Office Chair", **{"Warranty Years": 5})]))).json()["imported"] == 2
    first, second = await assets()
    assert (first.brand_id, first.warranty_years, first.description) == (w["brand"], 2, "Office Chair Dell")
    assert second.warranty_years == 5


async def test_an_item_outside_the_callers_work_area_is_a_row_error(client):
    w = await build_s4_world()
    content = xlsx([row(Item="Office Chair"), row(Item="Gang Box")])
    async with SessionLocal() as s:
        valid, errors, _ = await _validate_add_rows(s, content, None, ("IT",))
    assert [r["row"] for r in valid] == [3]
    assert errors == [{"row": 2, "field": "Item", "message": "The Item 'Office Chair' is Non-IT work, which is outside your work area."}]
